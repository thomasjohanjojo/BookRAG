import os
from typing import Optional
from docx import Document
from Interfaces.DocumentWriterToolAbstractBaseClass import DocumentWriterToolAbstractBaseClass


class DocxDocumentWriter(DocumentWriterToolAbstractBaseClass):
    """Concrete implementation for creating, appending, and saving notes to .docx files."""

    def write_note_to_document(
        self,
        content: str,
        file_path: str,
        heading: Optional[str] = None,
        append: bool = True
    ) -> str:
        try:
            # Ensure the directory exists if a path includes subdirectories
            folder = os.path.dirname(file_path)
            if folder and not os.path.exists(folder):
                os.makedirs(folder, exist_ok=True)

            # Open existing document to append, or start a new document
            if append and os.path.exists(file_path):
                doc = Document(file_path)
                status_verb = "Appended note to existing file"
            else:
                doc = Document()
                status_verb = "Created and wrote note to new file"

            # Add optional heading
            if heading:
                doc.add_heading(heading, level=2)

            # Append the body text
            doc.add_paragraph(content)

            # Save changes to disk
            doc.save(file_path)

            heading_info = f" under heading '{heading}'" if heading else ""
            return f"Success: {status_verb} '{file_path}'{heading_info}."

        except Exception as e:
            return f"Error: Failed to write to '{file_path}'. Reason: {str(e)}"